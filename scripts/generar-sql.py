"""Genera database/local/instalar.sql desde el Excel y el .env local.

El archivo de salida contiene nombres de solicitantes y la contraseña de MySQL.
No se sube a GitHub: database/local/ está en .gitignore.
"""

from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path
from urllib.parse import unquote, urlparse

import bcrypt
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
ENV_PATH = ROOT / ".env"
XLSX = ROOT / "BitacoraSoportes_Data.xlsx"
SCHEMA = (ROOT / "database" / "001-esquema.sql").read_text(encoding="utf-8")
OUT = ROOT / "database" / "local" / "instalar.sql"

TIPOS = {"Chequeo", "Instalación", "Correctivo"}


def load_env(path: Path) -> dict[str, str]:
    values: dict[str, str] = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        values[key.strip()] = value.strip()
    return values


def sql_str(value: object) -> str:
    if value is None:
        return "NULL"
    text = str(value).replace("\\", "\\\\").replace("'", "''")
    return f"'{text}'"


def sql_date(value: object) -> str:
    if value is None or value == "":
        return "NULL"
    if isinstance(value, datetime):
        return f"'{value.strftime('%Y-%m-%d')}'"
    text = str(value).strip()
    if len(text) >= 10 and text[4] == "-" and text[7] == "-":
        return f"'{text[:10]}'"
    for fmt in ("%d/%m/%Y", "%d/%m/%Y, %H:%M:%S"):
        try:
            return f"'{datetime.strptime(text[:19], fmt).strftime('%Y-%m-%d')}'"
        except ValueError:
            continue
    return "NULL"


def sql_datetime(value: object) -> str:
    if value is None or value == "":
        return "NULL"
    if isinstance(value, datetime):
        return f"'{value.strftime('%Y-%m-%d %H:%M:%S')}'"
    text = str(value).strip()
    for fmt in ("%d/%m/%Y, %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            parsed = datetime.strptime(text[:19] if "H" in fmt else text[:10], fmt)
            return f"'{parsed.strftime('%Y-%m-%d %H:%M:%S')}'"
        except ValueError:
            continue
    return "NULL"


def record_id(value: object) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, float):
        return str(int(value))
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    if not text.isdigit():
        return None
    return text


def main() -> None:
    env = load_env(ENV_PATH)
    url = urlparse(env["DATABASE_URL"])
    password = unquote(url.password or "")
    user = unquote(url.username or "")
    database = (url.path or "").lstrip("/")
    if user != "bitacora_admin" or database != "bitacora_campbell" or not password:
        raise SystemExit("DATABASE_URL debe apuntar a bitacora_admin y bitacora_campbell")

    admin_email = env["APP_ADMIN_EMAIL"].lower()
    admin_name = env["APP_ADMIN_NOMBRE"]
    admin_password = env["APP_ADMIN_PASSWORD"]
    password_hash = bcrypt.hashpw(admin_password.encode("utf-8"), bcrypt.gensalt(rounds=10)).decode("ascii")

    workbook = openpyxl.load_workbook(XLSX, data_only=True, read_only=True)
    helpers = []
    for (nombre,) in workbook["Helpers"].iter_rows(min_col=1, max_col=1, values_only=True):
        if nombre and str(nombre).strip():
            helpers.append(str(nombre).strip())

    rows = []
    extras = []
    tecnicos = []
    for raw in workbook["Data"].iter_rows(min_row=3, max_col=10, values_only=True):
        soporte_id = record_id(raw[0])
        if not soporte_id or soporte_id == "0":
            continue
        area = str(raw[2]).strip() if raw[2] else ""
        tecnico = str(raw[8]).strip() if raw[8] else ""
        if area and area not in helpers and area not in extras:
            extras.append(area)
        if tecnico and tecnico not in tecnicos:
            tecnicos.append(tecnico)
        rows.append(raw)

    areas = helpers + extras
    area_id = {nombre: index for index, nombre in enumerate(areas, start=1)}
    tecnico_id = {nombre: index for index, nombre in enumerate(tecnicos, start=1)}

    area_values = ",\n".join(f"  ({index}, {sql_str(nombre)})" for index, nombre in enumerate(areas, start=1))
    tecnico_values = ",\n".join(
        f"  ({index}, {sql_str(nombre)})" for index, nombre in enumerate(tecnicos, start=1)
    )

    soporte_rows = []
    for raw in rows:
        soporte_id = record_id(raw[0])
        area = str(raw[2]).strip() if raw[2] else ""
        tecnico = str(raw[8]).strip() if raw[8] else ""
        tipo = str(raw[4]).strip() if raw[4] else ""
        tipo_sql = sql_str(tipo) if tipo in TIPOS else "NULL"
        solicitante = str(raw[1]).strip() if raw[1] else "Sin nombre"
        descripcion = str(raw[3]).strip() if raw[3] else None
        tiempo = str(raw[7]).strip() if raw[7] else None
        actualizado = sql_datetime(raw[9])
        creado = actualizado if actualizado != "NULL" else sql_datetime(raw[5])
        if creado == "NULL":
            creado = "'2023-08-01 00:00:00'"
        soporte_rows.append(
            "("
            + ", ".join(
                [
                    sql_str(soporte_id),
                    sql_str(solicitante[:180]),
                    str(area_id[area]) if area in area_id else "NULL",
                    sql_str(descripcion) if descripcion else "NULL",
                    tipo_sql,
                    sql_date(raw[5]),
                    sql_date(raw[6]),
                    sql_str(tiempo[:80]) if tiempo else "NULL",
                    str(tecnico_id[tecnico]) if tecnico in tecnico_id else "NULL",
                    actualizado,
                    creado,
                    "1",
                ]
            )
            + ")"
        )

    lotes = []
    for start in range(0, len(soporte_rows), 100):
        chunk = ",\n".join(soporte_rows[start : start + 100])
        lotes.append(
            "INSERT IGNORE INTO soportes (\n"
            "  id, solicitante, area_id, descripcion, tipo, fecha_solicitud, fecha_entrega,\n"
            "  tiempo_soporte, tecnico_id, actualizado_en, creado_en, version\n"
            f") VALUES\n{chunk};"
        )

    quoted_password = password.replace("\\", "\\\\").replace("'", "''")
    header = f"""-- Carga local de Bitácora Campbell. Ejecutar como root en el puerto 3307.
-- Este archivo queda solo en esta PC.

CREATE DATABASE IF NOT EXISTS bitacora_campbell
  CHARACTER SET utf8mb4
  COLLATE utf8mb4_unicode_ci;

CREATE USER IF NOT EXISTS 'bitacora_admin'@'127.0.0.1' IDENTIFIED BY '{quoted_password}';
CREATE USER IF NOT EXISTS 'bitacora_admin'@'localhost' IDENTIFIED BY '{quoted_password}';
ALTER USER 'bitacora_admin'@'127.0.0.1' IDENTIFIED BY '{quoted_password}';
ALTER USER 'bitacora_admin'@'localhost' IDENTIFIED BY '{quoted_password}';
GRANT ALL PRIVILEGES ON bitacora_campbell.* TO 'bitacora_admin'@'127.0.0.1';
GRANT ALL PRIVILEGES ON bitacora_campbell.* TO 'bitacora_admin'@'localhost';
FLUSH PRIVILEGES;

"""

    catalogos = f"""
INSERT INTO areas (id, nombre) VALUES
{area_values}
ON DUPLICATE KEY UPDATE nombre = VALUES(nombre);

INSERT INTO tecnicos (id, nombre) VALUES
{tecnico_values}
ON DUPLICATE KEY UPDATE nombre = VALUES(nombre);

INSERT INTO usuarios (id, nombre, email, password_hash, rol, activo)
VALUES (1, {sql_str(admin_name)}, {sql_str(admin_email)}, {sql_str(password_hash)}, 'admin', 1)
ON DUPLICATE KEY UPDATE
  nombre = VALUES(nombre),
  password_hash = VALUES(password_hash),
  rol = 'admin',
  activo = 1;
"""

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(header + SCHEMA + "\n" + catalogos + "\n" + "\n\n".join(lotes) + "\n", encoding="utf-8")
    print(f"areas={len(areas)} tecnicos={len(tecnicos)} soportes={len(soporte_rows)} archivo={OUT}")


if __name__ == "__main__":
    main()
